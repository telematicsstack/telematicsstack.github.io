#!/usr/bin/env python3
"""Build decks from YAML content files using the committed .potx template.

Usage:
    uv run build.py content/overview.yaml [content/more.yaml ...]
    uv run build.py            # builds every YAML in content/

Reads DECK_VERSION from the environment (default "dev") and puts it on the
closing slide's contact line. Writes decks/dist/<deck-name>.pptx and, when
LibreOffice (soffice) is on PATH, a PDF next to it.

Slides are added by layout name; placeholders are filled by shape name (with
the placeholder idx as tie-breaker), so the YAML never positions anything and
no shapes are drawn that the template already provides.
"""

from __future__ import annotations

import io
import json
import os
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import yaml
from pptx import Presentation
from pptx.util import Pt

import prepare

DECKS = Path(__file__).resolve().parent
TEMPLATE = DECKS / "template" / "telematics-stack.potx"
CONTENT = DECKS / "content"
DIST = DECKS / "dist"
TOKENS = json.loads((DECKS.parent / "design-system" / "tokens.json").read_text())


def token(name: str) -> str:
    for t in TOKENS["color"]["tokens"]:
        if t["name"] == name:
            return t["value"].lstrip("#").upper()
    raise KeyError(f"colour token not found: {name}")


def open_potx(path: Path) -> Presentation:
    """python-pptx refuses the template content type, so patch it in memory."""
    src = zipfile.ZipFile(path)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"presentationml.template.main+xml",
                    b"presentationml.presentation.main+xml",
                )
            dst.writestr(item, data)
    buf.seek(0)
    return Presentation(buf)


def find_layout(prs: Presentation, name: str):
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            if layout.name == name:
                return layout
    raise KeyError(f"layout not found in template: {name!r}")


def ph_by_name(slide, name: str):
    """python-pptx renames cloned placeholders, but keeps their idx — so the
    template's placeholder names are resolved to an idx via the slide's layout."""
    for layout_ph in slide.slide_layout.placeholders:
        if layout_ph.name == name:
            idx = layout_ph.placeholder_format.idx
            for shape in slide.placeholders:
                if shape.placeholder_format.idx == idx:
                    return shape
    raise KeyError(
        f"placeholder {name!r} not on layout {slide.slide_layout.name!r} "
        f"(has: {[s.name for s in slide.slide_layout.placeholders]})"
    )


def set_text(slide, name: str, value: str, upper: bool = False) -> None:
    value = str(value)
    if upper:
        value = value.upper()
    ph_by_name(slide, name).text_frame.text = value


def set_cell_borders(cell, bottom_hex: str | None) -> None:
    """Hairline (1px line-soft) rule under the row, no other borders — the
    brand separates rows, never columns. python-pptx has no border API, and a
    styleless table would otherwise get renderer-default borders."""
    from pptx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for el in tc_pr.findall(qn(tag)):
            tc_pr.remove(el)

    def line(tag: str, hex_val: str | None):
        ln = tc_pr.makeelement(qn(tag), {"w": "9525", "cap": "flat"})  # 1px = 9525 EMU
        if hex_val:
            fill = ln.makeelement(qn("a:solidFill"), {})
            clr = ln.makeelement(qn("a:srgbClr"), {"val": hex_val})
            fill.append(clr)
            ln.append(fill)
        else:
            ln.append(ln.makeelement(qn("a:noFill"), {}))
        return ln

    # border elements must precede any fill element inside tcPr, in L/R/T/B order
    for pos, (tag, hex_val) in enumerate(
        (("a:lnL", None), ("a:lnR", None), ("a:lnT", None), ("a:lnB", bottom_hex))
    ):
        tc_pr.insert(pos, line(tag, hex_val))


def fill_table(slide, rows: list[list[str]]) -> None:
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn

    ph = ph_by_name(slide, "table")
    n_rows, n_cols = len(rows), max(len(r) for r in rows)
    frame = ph.insert_table(rows=n_rows, cols=n_cols)
    table = frame.table
    table.first_row = False  # no header banding — every cell is a spec cell
    table.horz_banding = False
    line_soft = token("line-soft")
    ink = token("ink")
    ink_muted = token("ink-muted")
    for r, row in enumerate(rows):
        table.rows[r].height = Pt(30)
        for c in range(n_cols):
            cell = table.cell(r, c)
            cell.fill.background()  # ground shows through; no zebra fills
            cell.text = str(row[c]) if c < len(row) else ""
            for para in cell.text_frame.paragraphs:
                for run in para.runs or [para.add_run()]:
                    run.font.name = "Space Mono"
                    run.font.size = Pt(14)
                    run.font.bold = c == 0  # mono-bold label column
                    run.font.color.rgb = RGBColor.from_string(ink if c == 0 else ink_muted)
            set_cell_borders(cell, line_soft if r < n_rows - 1 else None)
    # first column narrow, rest split evenly
    if n_cols > 1:
        total = sum(col.width for col in table.columns)
        first = int(total * 0.25)
        rest = (total - first) // (n_cols - 1)
        table.columns[0].width = first
        for c in range(1, n_cols):
            table.columns[c].width = rest
    # drop the default table style so no theme banding colours leak in
    tbl = frame.table._tbl
    tbl_pr = tbl.find(qn("a:tblPr"))
    if tbl_pr is not None:
        for el in tbl_pr.findall(qn("a:tableStyleId")):
            tbl_pr.remove(el)


FILLERS = {}


def filler(layout_name):
    def wrap(fn):
        FILLERS[layout_name] = fn
        return fn
    return wrap


@filler("Title")
def fill_title(slide, entry):
    set_text(slide, "eyebrow", entry["eyebrow"], upper=True)
    set_text(slide, "title", entry["title"], upper=True)
    set_text(slide, "subtitle", entry["subtitle"])


@filler("Section")
def fill_section(slide, entry):
    set_text(slide, "eyebrow", entry["eyebrow"], upper=True)
    set_text(slide, "title", entry["title"], upper=True)


@filler("Stat row")
def fill_stat_row(slide, entry):
    set_text(slide, "title", entry["title"], upper=True)
    stats = entry["stats"]
    if len(stats) != 4:
        raise ValueError(f"Stat row needs exactly 4 stats, got {len(stats)}")
    for i, stat in enumerate(stats, start=1):
        set_text(slide, f"stat{i}-value", stat["value"])
        set_text(slide, f"stat{i}-label", stat["label"], upper=True)


@filler("Feature cards")
def fill_feature_cards(slide, entry):
    set_text(slide, "title", entry["title"], upper=True)
    cards = entry["cards"]
    if len(cards) != 3:
        raise ValueError(f"Feature cards needs exactly 3 cards, got {len(cards)}")
    for i, card in enumerate(cards, start=1):
        set_text(slide, f"card{i}-eyebrow", card["eyebrow"], upper=True)
        set_text(slide, f"card{i}-heading", card["heading"])
        set_text(slide, f"card{i}-body", card["body"])
        set_text(slide, f"card{i}-spec", card["spec"])


@filler("Spec table")
def fill_spec_table(slide, entry):
    set_text(slide, "title", entry["title"], upper=True)
    fill_table(slide, entry["table"])


@filler("Closing")
def fill_closing(slide, entry, version="dev"):
    set_text(slide, "title", entry["title"], upper=True)
    set_text(slide, "contact", f"{entry['contact']} · {version}")


def build(content_path: Path, version: str) -> Path:
    spec = yaml.safe_load(content_path.read_text())
    prs = open_potx(TEMPLATE)
    for entry in spec["slides"]:
        layout = find_layout(prs, entry["layout"])
        slide = prs.slides.add_slide(layout)
        fn = FILLERS[entry["layout"]]
        if entry["layout"] == "Closing":
            fn(slide, entry, version=version)
        else:
            fn(slide, entry)
    DIST.mkdir(exist_ok=True)
    out = DIST / f"{content_path.stem}.pptx"
    prs.save(out)
    print(f"Built {out.relative_to(DECKS)} ({len(spec['slides'])} slides, version {version})")
    return out


def to_pdf(pptx_path: Path) -> None:
    soffice = shutil.which("soffice")
    if not soffice:
        print(f"  note: soffice not found — skipping PDF for {pptx_path.name}")
        return
    subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(DIST), str(pptx_path)],
        check=True,
        capture_output=True,
    )
    print(f"  PDF: {pptx_path.with_suffix('.pdf').relative_to(DECKS)} (LibreOffice preview render)")


def main() -> None:
    prepare.main()  # derive logos/fonts into .cache/ so a fresh clone works
    version = os.environ.get("DECK_VERSION", "dev")
    paths = [Path(a) for a in sys.argv[1:]] or sorted(CONTENT.glob("*.yaml"))
    if not paths:
        sys.exit("no content files found in decks/content/")
    for path in paths:
        pptx_path = build(path, version)
        to_pdf(pptx_path)


if __name__ == "__main__":
    main()
